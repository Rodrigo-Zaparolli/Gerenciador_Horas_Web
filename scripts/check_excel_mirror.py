"""Render synthetic spreadsheets with Excel; never open user workbooks."""
import base64
import hashlib
import io
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.formatting.rule import DataBarRule
from PIL import Image
from local_edesk.excel_mirror import render

output = Path(__file__).resolve().parents[1] / 'build/excel-mirror-validation'
output.mkdir(parents=True, exist_ok=True)
for index, name in enumerate(['Trabalho A', 'Outro trabalho']):
    book = Workbook()
    sheet = book.active
    sheet.title = name
    sheet.merge_cells('A1:D1')
    sheet['A1'] = name
    sheet['A1'].fill = PatternFill('solid', fgColor='60C8E8' if index == 0 else 'F2BE40')
    sheet['A1'].font = Font(bold=True, size=14)
    sheet['A1'].alignment = Alignment(horizontal='center')
    for column, text in enumerate(['Etapa', 'Previsto', 'Realizado', 'Percentual'], 1):
        cell = sheet.cell(2, column, text)
        cell.fill = PatternFill('solid', fgColor='712477')
        cell.font = Font(color='FFFFFF', bold=True)
        sheet.column_dimensions[cell.column_letter].width = 24 + index * 4
    sheet['A3'], sheet['B3'], sheet['C3'], sheet['D3'] = 'Implantação', 9/24, 2/24, .25
    sheet['B3'].number_format = sheet['C3'].number_format = '[h]:mm'
    sheet['D3'].number_format = '0%'
    sheet.conditional_formatting.add('D3', DataBarRule(start_type='num',start_value=0,end_type='num',end_value=1,color='63C384'))
    sheet.row_dimensions[1].height = 28
    sheet.row_dimensions[2].height = 26
    path = output / f'fixture-{index}.xlsx'
    book.save(path)
    previous = hashlib.sha256(path.read_bytes()).digest()
    result = render(path, f"'{name}'!A1:D3")
    assert result['renderer'] == 'excel'
    assert hashlib.sha256(path.read_bytes()).digest() == previous
    png = base64.b64decode(result['imagemBase64'])
    image = Image.open(io.BytesIO(png)).convert('RGB')
    assert image.width > 400 and image.height > 50, image.size
    purple = sum(1 for r,g,b in image.getdata() if r > 60 and b > 70 and g < 70)
    assert purple > 100, 'Original purple header missing'
    (output / f'fixture-{index}.png').write_bytes(png)
    print(f'workbook={index}; native Excel image={image.width}x{image.height}; source unchanged; original header retained')
