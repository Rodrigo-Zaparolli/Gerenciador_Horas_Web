"""Real Excel acceptance on a generated workbook only; never edit user files."""
import sys
from pathlib import Path
from datetime import datetime
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from local_edesk.excel_editor import edit

folder = ROOT / 'build/workbook-sheet-validation'
folder.mkdir(parents=True, exist_ok=True)
path = folder / 'synthetic-worksheet.xlsx'
book = Workbook()
sheet = book.active
sheet.title = 'Qualquer guia'
sheet['A1'] = 'Antes'
sheet['B1'] = datetime(2026, 10, 23)
sheet['B1'].number_format = 'dd/mm/yyyy'
sheet['C1'] = '=1+1'
book.save(path)
args = {'worksheet': sheet.title, 'firstRow': 1, 'firstColumn': 1}
before = edit(path, {'action': 'excel-editor-read', **args})
fields = {field['cell']: field for field in before['fields']}
assert fields['C1']['readOnly'] is True
changes = [{'sheet': sheet.title, 'cell': address, 'before': fields[address]['value'], 'value': value}
           for address, value in [('A1', 'Depois'), ('B1', '2026-10-24'), ('D2', 12.5)]]
saved = edit(path, {'action': 'excel-editor-write', **args,
                    'revision': before['revision'], 'changes': changes})
assert saved['saved'] is True
checked = load_workbook(path)
assert checked[sheet.title]['A1'].value == 'Depois'
assert checked[sheet.title]['B1'].value.date().isoformat() == '2026-10-24'
assert checked[sheet.title]['C1'].value == '=1+1'
assert checked[sheet.title]['D2'].value == 12.5
assert Path(saved['backup']).exists()
print('WORKBOOK_SHEET_REAL_EXCEL_OK: text/date/new numeric cell, formula preserved, backup created')
